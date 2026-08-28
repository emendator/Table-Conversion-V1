"""
表格转换工具 - FastAPI 后端服务（React + FastAPI 架构）
提供纯 REST API，不渲染模板
"""

import os
import sys
import time
import uuid
import base64
import hashlib
import secrets
import threading
from typing import Optional
from contextlib import asynccontextmanager

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request, Depends
from fastapi.responses import JSONResponse, FileResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

import ai_config
import auth
import user_system

# 添加当前目录到路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from config import (
    BACKEND_DIR, OUTPUT_DIR, UPLOAD_DIR,
    IMAGE_ROW_HEIGHT_PT, IMAGE_MARGIN_PT,
    DEEPSEEK_MODEL, DEEPSEEK_API_URL,
)
from table_converter import (
    read_source_file, read_reference_template,
    normalize_header, get_formula_headers, extract_images_from_source,
    detect_source_formulas,
)
from deepseek_matcher import match_headers
from mapping_history import (
    save_convert_mappings,
    list_mapping_records, add_mapping_record,
    update_mapping_record, delete_mapping_record, clear_mapping_records,
    import_mapping_records, batch_delete_mapping_records,
    check_duplicates, preview_convert_duplicates,
    migrate_company_data,
)

# ============================================================
# 缓存定时清理（uploads / 输出）
# ============================================================
CLEANUP_INTERVAL = 5 * 60   # 每 5 分钟清理一次，防止缓存堆积
CLEANUP_MAX_AGE = 1 * 3600  # 清理超过 1 小时的临时文件

_cleanup_stop = threading.Event()


def _cleanup_worker():
    """后台定时清理线程：每 5 分钟清理 uploads 与 输出 目录的过期缓存"""
    while not _cleanup_stop.wait(CLEANUP_INTERVAL):
        try:
            cleanup_old_files()
        except Exception as e:
            print(f"  ⚠️ 定时清理异常: {e}")


@asynccontextmanager
async def lifespan(_app: FastAPI):
    """应用生命周期：启动时开启缓存清理后台线程，退出时停止"""
    _cleanup_stop.clear()
    t = threading.Thread(target=_cleanup_worker, daemon=True, name="cache-cleaner")
    t.start()
    print("  🧹 缓存定时清理已启动（每 5 分钟）")
    try:
        yield
    finally:
        _cleanup_stop.set()


# ============================================================
# FastAPI 应用初始化
# ============================================================
app = FastAPI(
    title="表格转换工具 API",
    version="3.0",
    description="将客户发来的原始表格自动转换为公司内部参考表格格式",
    lifespan=lifespan,
)

# CORS 配置（前后端分离必需）
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 确保目录存在
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

# 初始化内置最高级管理员（admin / ***REMOVED***，公司：***REMOVED***）
user_system.init_super_admin()

# ============================================================
# 会话管理（基于 X-Session-Id 请求头）
# ============================================================
sessions = {}


def get_session_id(request: Request) -> str:
    sid = request.headers.get("X-Session-Id", "").strip()
    if not sid:
        sid = uuid.uuid4().hex
    return sid


def get_session(request: Request) -> dict:
    sid = get_session_id(request)
    if sid not in sessions:
        sessions[sid] = {
            "source_paths": [],
            "source_files": [],
            "reference_path": None,
            "reference_headers": [],
            "mappings_list": [],
            "output_path": None,
            "original_mappings": [],
        }
    return sessions[sid]


def cleanup_old_files():
    now = time.time()
    cleaned = 0
    for directory in (UPLOAD_DIR, OUTPUT_DIR):
        if not os.path.isdir(directory):
            continue
        for fname in os.listdir(directory):
            fpath = os.path.join(directory, fname)
            try:
                if os.path.isfile(fpath) and now - os.path.getmtime(fpath) > CLEANUP_MAX_AGE:
                    os.remove(fpath)
                    cleaned += 1
            except Exception:
                pass
    if cleaned:
        print(f"  🧹 已清理 {cleaned} 个过期临时文件")


def _cleanup_session_files(request: Request) -> None:
    """删除当前会话产生的上传/输出临时文件并移除会话（游客退出登录时清除使用痕迹用）"""
    sid = get_session_id(request)
    data = sessions.pop(sid, None)
    if not data:
        return
    paths = list(data.get("source_paths", []))
    if data.get("reference_path"):
        paths.append(data["reference_path"])
    if data.get("output_path"):
        paths.append(data["output_path"])
    for p in paths:
        try:
            if p and os.path.isfile(p):
                os.remove(p)
        except Exception:
            pass


# ============================================================
# API Key 传输加密（前端↔后端）
# 前端使用会话 nonce 对 Key 做 XOR 流加密后上传，避免明文出现在网络请求中。
# ============================================================
TRANSPORT_SALT = "table-converter-ai-transport-v1"


def _transport_decrypt(token: str, nonce: str) -> str:
    """解密前端用 nonce 加密的 API Key（与前端 encryptWithNonce 对应）"""
    if not token or not nonce:
        return ""
    key = hashlib.sha256((TRANSPORT_SALT + nonce).encode("utf-8")).digest()
    b64 = token.replace('-', '+').replace('_', '/')
    b64 += '=' * (-len(b64) % 4)
    data = base64.b64decode(b64)
    return bytes(b ^ key[i % len(key)] for i, b in enumerate(data)).decode("utf-8")


# ============================================================
# Pydantic 模型
# ============================================================
class MatchRequest(BaseModel):
    useAI: bool = True


class MappingItem(BaseModel):
    source: str
    target: str = ""
    confidence: Optional[str] = "manual"
    reason: Optional[str] = ""
    add_to_library: bool = False  # 用户是否勾选「加入映射库」
    # 公式列相关：
    source_is_formula: bool = False   # 源列是否为原始表中的公式列
    target_is_formula: bool = False   # 目标列是否为参考表中的公式列
    # 公式处理方式：
    #   default        无公式，正常写源数据
    #   target_formula 目标列是参考公式 → 保留参考公式
    #   target_data    目标列是参考公式 → 用源数据覆盖
    #   source_formula 源列是源公式 → 写源公式
    #   source_data    源列是源公式 → 写源数据（计算值）
    formula_action: str = "default"


class ConvertRequest(BaseModel):
    mappings: list[MappingItem] = []
    on_duplicate: str = "overwrite"  # 保存映射到库时重复处理: overwrite|skip|abort


class DeleteSourceRequest(BaseModel):
    """删除单个原始表格文件请求"""
    filename: str = ""


class MappingLibItem(BaseModel):
    """映射库新增/更新请求"""
    source: str = ""
    target: str = ""
    on_duplicate: str = "overwrite"  # 重复处理: overwrite|skip|abort


class CheckDuplicateRequest(BaseModel):
    """查重请求"""
    records: list[MappingLibItem] = []


class ImportRequest(BaseModel):
    """映射库导入请求"""
    records: list[MappingLibItem] = []
    mode: str = "merge"  # merge 合并 | replace 覆盖
    on_duplicate: str = "overwrite"  # merge 模式下重复处理: overwrite|skip|abort


class BatchDeleteRequest(BaseModel):
    """映射库批量删除请求"""
    ids: list[str] = []


class AIConfigRequest(BaseModel):
    """AI API 配置保存请求"""
    api_key_enc: str = ""   # 前端用会话 nonce 加密后的 API Key（推荐）
    api_key: str = ""       # 兜底：明文 Key（本地使用，传输中仍建议加密）


# ============================================================
# 用户系统请求模型（注册 / 登录 / 邀请码 / 用户管理）
# ============================================================

class RegisterRequest(BaseModel):
    """注册请求：用户名 + 密码 + 邮箱 + 邀请码 + 公司名称 + 公司数据库密码

    公司数据库密码：公司已存在则必须填写并校验匹配才能加入；
    公司不存在则以此密码创建新公司数据库。
    """
    username: str = ""
    password: str = ""
    email: str = ""
    invite_code: str = ""
    company: str = ""
    company_db_password: str = ""


class LoginRequest(BaseModel):
    """登录请求"""
    username: str = ""
    password: str = ""


class InviteGenerateRequest(BaseModel):
    """生成邀请码请求"""
    level: int = 2     # 1=管理员(一级) 2=用户(二级)
    expires_at: str = ""   # 账户有效期截止日期（空 = 永久）


class InviteExtendRequest(BaseModel):
    """延长邀请码请求"""
    code: str = ""
    expires_at: str = ""   # 新的有效期截止日期（空 = 永久）


class InviteCodeRequest(BaseModel):
    """删除邀请码请求"""
    code: str = ""


class UserExtendRequest(BaseModel):
    """延长账号请求"""
    expires_at: str = ""   # 新的有效期截止日期（空 = 永久）


class CompanyPasswordRequest(BaseModel):
    """创建/修改公司数据库密码请求"""
    new_password: str = ""
    company: str = ""   # 仅最高级管理员可指定其他公司；留空为当前用户所在公司



# ============================================================
# API 路由
# ============================================================

@app.get("/")
async def root():
    return {
        "service": "表格转换工具 API",
        "version": "3.0",
        "architecture": "React + FastAPI",
        "endpoints": {
            "upload_source": "POST /api/upload/source",
            "upload_reference": "POST /api/upload/reference",
            "match": "POST /api/match",
            "convert": "POST /api/convert",
            "download": "GET /api/download",
            "reset": "POST /api/reset",
            "clear_source": "POST /api/clear/source",
            "clear_reference": "POST /api/clear/reference",
            "mapping_library_list": "GET /api/mapping-library",
            "mapping_library_add": "POST /api/mapping-library",
            "mapping_library_update": "PUT /api/mapping-library/{id}",
            "mapping_library_delete": "DELETE /api/mapping-library/{id}",
            "mapping_library_clear": "DELETE /api/mapping-library",
        },
    }


# ============================================================
# 用户系统：认证（注册 / 登录 / 登出 / 当前用户）
# ============================================================

@app.post("/api/auth/register")
async def api_register(body: RegisterRequest):
    """注册：用户名 + 密码 + 邮箱 + 邀请码 + 公司名称 + 公司数据库密码"""
    result = user_system.register(
        body.username, body.password, body.email,
        body.invite_code, body.company, body.company_db_password,
    )
    if not result["success"]:
        raise HTTPException(status_code=400, detail=result["message"])
    token = auth.create_token(result["user"]["username"])
    return {"success": True, "message": result["message"], "user": result["user"], "token": token}


@app.post("/api/auth/login")
async def api_login(body: LoginRequest):
    """登录：用户名 + 密码"""
    result = user_system.login(body.username, body.password)
    if not result["success"]:
        # 结构化错误：detail 为 {message, code}，供前端识别「账户已过期」等特殊场景
        raise HTTPException(status_code=403, detail={
            "message": result["message"],
            "code": result.get("code", ""),
        })
    token = auth.create_token(result["user"]["username"])
    return {"success": True, "message": result["message"], "user": result["user"], "token": token}


@app.post("/api/auth/logout")
async def api_logout(request: Request):
    """登出：使当前令牌失效；游客登出时自动清除全部使用痕迹"""
    token = auth._extract_token(request)
    user = auth.get_user_by_token(token)
    auth.revoke_token(token)
    if user and user.get("is_guest"):
        # 游客：删除账号 + 映射库/API Key + 本次会话产生的上传/输出文件
        user_system.delete_guest(user.get("username", ""))
        _cleanup_session_files(request)
    return {"success": True}


@app.post("/api/auth/guest")
async def api_guest():
    """游客登录：创建临时游客账号（独立映射库，退出登录后自动清除全部使用痕迹）"""
    user = user_system.create_guest()
    token = auth.create_token(user["username"])
    return {"success": True, "message": "游客登录成功", "user": user, "token": token}


@app.get("/api/auth/me")
async def api_me(user: dict = Depends(auth.require_login)):
    """当前登录用户信息"""
    return {"success": True, "user": user}


# ============================================================
# 用户系统：邀请码（admin 可一级+二级；manager 可二级）
# ============================================================

@app.post("/api/invite/generate")
async def api_invite_generate(body: InviteGenerateRequest,
                              user: dict = Depends(auth.require_manager)):
    """生成邀请码（level: 1=管理员 2=用户）"""
    result = user_system.generate_invite(user["username"], body.level, body.expires_at)
    if not result["success"]:
        raise HTTPException(status_code=400, detail=result["message"])
    return result


@app.get("/api/invite/list")
async def api_invite_list(user: dict = Depends(auth.require_manager)):
    """查看邀请码列表"""
    return {"success": True, "invites": user_system.list_invites(user["username"])}


@app.post("/api/invite/validate")
async def api_invite_validate(body: InviteCodeRequest):
    """校验邀请码（只读，不修改状态）：返回 valid / level / reason（供注册页识别邀请码级别）"""
    return {"success": True, **user_system.validate_invite(body.code)}


@app.post("/api/invite/extend")
async def api_invite_extend(body: InviteExtendRequest,
                            user: dict = Depends(auth.require_manager)):
    """延长邀请码有效时间"""
    result = user_system.extend_invite(user["username"], body.code, body.expires_at)
    if not result["success"]:
        raise HTTPException(status_code=400, detail=result["message"])
    return result


@app.post("/api/invite/delete")
async def api_invite_delete(body: InviteCodeRequest,
                            user: dict = Depends(auth.require_manager)):
    """删除邀请码（admin 任意；manager 仅自己创建的）"""
    result = user_system.delete_invite(user["username"], body.code)
    if not result["success"]:
        raise HTTPException(status_code=400, detail=result["message"])
    return result


# ============================================================
# 用户系统：用户管理（admin / manager，经理仅本公司）
# ============================================================

@app.get("/api/admin/users")
async def api_admin_users(user: dict = Depends(auth.require_manager)):
    """用户列表（admin 全量；manager 仅本公司）"""
    return {"success": True, "users": user_system.list_users(user["username"])}


@app.post("/api/admin/users/{user_id}/extend")
async def api_admin_user_extend(user_id: str, body: UserExtendRequest,
                                user: dict = Depends(auth.require_manager)):
    """延长账号有效时间"""
    result = user_system.extend_user(user["username"], user_id, body.expires_at)
    if not result["success"]:
        raise HTTPException(status_code=400, detail=result["message"])
    return result


@app.post("/api/admin/users/{user_id}/disable")
async def api_admin_user_disable(user_id: str,
                                 user: dict = Depends(auth.require_manager)):
    """清除账号（admin 清除 manager 时级联清除其下用户）"""
    result = user_system.disable_user(user["username"], user_id)
    if not result["success"]:
        raise HTTPException(status_code=400, detail=result["message"])
    return result


# ============================================================
# 公司数据库密码（映射库隔离凭证）管理
# ============================================================

@app.get("/api/company/info")
async def api_company_info(user: dict = Depends(auth.require_login)):
    """查看当前公司数据库信息（不返回明文密码）"""
    info = user_system.get_company_db_info(user["company"])
    if info is None:
        return {"success": True, "company": user["company"], "is_set": False}
    return {
        "success": True,
        "company": info["company"],
        "is_set": True,
        "created_at": info.get("created_at", ""),
        "updated_at": info.get("updated_at", ""),
        "created_by": info.get("created_by", ""),
    }


@app.post("/api/company/password")
async def api_company_password(body: CompanyPasswordRequest,
                               user: dict = Depends(auth.require_manager)):
    """创建/修改公司数据库密码（manager 仅本公司；admin 可指定任意公司）"""
    result = user_system.update_company_db_password(
        user["username"], body.company or user["company"], body.new_password
    )
    if not result["success"]:
        raise HTTPException(status_code=400, detail=result["message"])
    # 隔离键变更后，将旧键下的映射库文件迁移到新键，避免数据丢失
    migrate_company_data(result.get("old_key", ""), result.get("new_key", ""))
    return result


@app.post("/api/upload/source")
def upload_source(request: Request, files: list[UploadFile] = File(...), user: dict = Depends(auth.require_login)):
    """上传原始表格文件（支持多文件）

    同步 def 路由：FastAPI 会将其放入线程池执行，避免 pandas/openpyxl 解析大文件时阻塞事件循环。
    """
    cleanup_old_files()
    data = get_session(request)
    uploaded = []
    errors = []

    for file in files:
        if file.filename == "" or file.filename is None:
            continue
        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in (".xlsx", ".xls"):
            errors.append(f"{file.filename}: 仅支持 .xlsx / .xls 格式")
            continue

        fname = f"source_{uuid.uuid4().hex[:8]}{ext}"
        filepath = os.path.join(UPLOAD_DIR, fname)
        content = file.file.read()
        os.makedirs(UPLOAD_DIR, exist_ok=True)   # 上传目录被删除后自动重建
        with open(filepath, "wb") as f:
            f.write(content)

        try:
            df, headers, source_header_row = read_source_file(filepath)
            preview = df.head(3).fillna("").to_dict(orient="records")
            for row in preview:
                for k, v in row.items():
                    if isinstance(v, float):
                        row[k] = round(v, 4)

            # 准确识别原始表格中的公式列（无硬编码）
            source_formula_info = detect_source_formulas(filepath, header_row=source_header_row)
            source_formula_columns = []
            for col_idx in range(1, len(headers) + 1):
                from openpyxl.utils import get_column_letter as _gcl
                col_letter = _gcl(col_idx)
                if col_letter in source_formula_info:
                    source_formula_columns.append(headers[col_idx - 1])

            file_info = {
                "name": os.path.splitext(file.filename)[0],
                "filename": file.filename,
                "path": filepath,
                "headers": headers,
                "rows": len(df),
                "header_row": source_header_row,
                "preview": preview,
                "source_formula_info": source_formula_info,      # 源表公式列: 列字母→公式
                "source_formula_columns": source_formula_columns,  # 源表公式列名
            }

            # 同名文件去重：同一会话重复上传（拖拽/重复选择同一文件）
            # 以最新一次为准——先移除旧记录与旧文件，避免列表里出现两个相同文件
            existing_idx = next(
                (i for i, f in enumerate(data["source_files"])
                 if f["filename"] == file.filename),
                None,
            )
            if existing_idx is not None:
                old = data["source_files"].pop(existing_idx)
                try:
                    os.remove(old["path"])
                except Exception:
                    pass
                if existing_idx < len(data["source_paths"]):
                    data["source_paths"].pop(existing_idx)

            data["source_files"].append(file_info)
            data["source_paths"].append(filepath)
            uploaded.append(file_info)
        except Exception as e:
            errors.append(f"{file.filename}: 解析失败 - {str(e)}")

    if not uploaded and errors:
        raise HTTPException(status_code=400, detail="; ".join(errors))

    return {
        "success": True,
        "sessionId": get_session_id(request),
        "files": [
            {"name": f["name"], "filename": f["filename"],
             "rows": f["rows"], "headers": f["headers"]}
            for f in data["source_files"]
        ],
        "totalFiles": len(data["source_files"]),
        "newFiles": [
            {"name": f["name"], "filename": f["filename"],
             "rows": f["rows"], "headers": f["headers"]}
            for f in uploaded
        ],
        "errors": errors if errors else None,
    }


@app.post("/api/upload/reference")
def upload_reference(request: Request, file: UploadFile = File(...), user: dict = Depends(auth.require_login)):
    """上传参考模板文件"""
    cleanup_old_files()
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名为空")

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in (".xlsx", ".xls"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xls 格式")

    fname = f"ref_{uuid.uuid4().hex[:8]}{ext}"
    filepath = os.path.join(UPLOAD_DIR, fname)
    content = file.file.read()
    os.makedirs(UPLOAD_DIR, exist_ok=True)   # 上传目录被删除后自动重建
    with open(filepath, "wb") as f:
        f.write(content)

    try:
        template = read_reference_template(filepath)
        data = get_session(request)
        data["reference_path"] = filepath
        data["reference_headers"] = template["headers"]
        data["formula_count"] = len(template["formula_info"])
        data["formula_info"] = {k: v for k, v in template["formula_info"].items()}

        return {
            "success": True,
            "sessionId": get_session_id(request),
            "filename": file.filename,
            "headers": template["headers"],
            "formulaColumns": list(template["formula_info"].keys()),
            "formulaCount": len(template["formula_info"]),
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"模板解析失败: {str(e)}")


@app.post("/api/clear/source")
async def clear_source(request: Request, user: dict = Depends(auth.require_login)):
    data = get_session(request)
    for src in data.get("source_files", []):
        try:
            os.remove(src["path"])
        except Exception:
            pass
    for p in data.get("source_paths", []):
        try:
            os.remove(p)
        except Exception:
            pass
    data["source_files"] = []
    data["source_paths"] = []
    data["mappings_list"] = []
    return {"success": True}


@app.post("/api/delete/source")
async def delete_source(request: Request, body: DeleteSourceRequest, user: dict = Depends(auth.require_login)):
    """删除单个原始表格文件"""
    data = get_session(request)
    filename = body.filename
    remaining = []
    deleted = False
    for src in data.get("source_files", []):
        if src.get("filename") == filename:
            try:
                os.remove(src["path"])
            except Exception:
                pass
            deleted = True
        else:
            remaining.append(src)
    if not deleted:
        raise HTTPException(status_code=404, detail="未找到该文件")
    data["source_files"] = remaining
    data["source_paths"] = [s.get("path", "") for s in remaining]
    data["mappings_list"] = []
    return {
        "success": True,
        "files": [
            {"name": f["name"], "filename": f["filename"], "rows": f["rows"], "headers": f["headers"]}
            for f in remaining
        ],
    }


@app.post("/api/clear/reference")
async def clear_reference(request: Request, user: dict = Depends(auth.require_login)):
    data = get_session(request)
    if data.get("reference_path") and os.path.exists(data["reference_path"]):
        try:
            os.remove(data["reference_path"])
        except Exception:
            pass
    data["reference_path"] = None
    data["reference_headers"] = []
    data["mappings_list"] = []
    return {"success": True}


# ============================================================
# 映射匹配库 CRUD API（前端数据管理页面）
# ============================================================

@app.get("/api/mapping-library")
async def api_mapping_library_list(user: dict = Depends(auth.require_login)):
    """获取映射库中所有记录"""
    records = list_mapping_records()
    return {"success": True, "records": records, "count": len(records)}


@app.post("/api/mapping-library")
async def api_mapping_library_add(body: MappingLibItem, user: dict = Depends(auth.require_login)):
    """新增一条映射记录"""
    if not body.source.strip():
        raise HTTPException(status_code=400, detail="源列名不能为空")
    if not body.target.strip():
        raise HTTPException(status_code=400, detail="目标列名不能为空")
    result = add_mapping_record(body.source, body.target,
                                on_duplicate=body.on_duplicate)
    return {"success": True, "action": result["action"], "record": result["record"]}


@app.post("/api/mapping-library/check-duplicates")
async def api_mapping_library_check_duplicates(body: CheckDuplicateRequest, user: dict = Depends(auth.require_login)):
    """检查一批待写入记录中，哪些源列名已存在于库中（仅以原始列名为准）"""
    records = [
        {"source": r.source, "target": r.target}
        for r in body.records
    ]
    duplicates = check_duplicates(records)
    return {"success": True, "duplicates": duplicates, "count": len(duplicates)}


@app.post("/api/mapping-library/check-convert-duplicates")
async def api_mapping_library_check_convert_duplicates(
        request: Request, body: ConvertRequest = ConvertRequest(),
        user: dict = Depends(auth.require_login)):
    """预检转换完成后将保存到库的映射中的重复项"""
    data = get_session(request)
    original_mappings = data.get("original_mappings", [])
    result = preview_convert_duplicates(
        [
            {
                "source": m.source,
                "target": m.target,
                "add_to_library": m.add_to_library,
            }
            for m in body.mappings
        ],
        original_mappings=original_mappings,
    )
    return {
        "success": True,
        "saveableCount": result["saveable_count"],
        "duplicates": result["duplicates"],
        "count": len(result["duplicates"]),
    }


@app.put("/api/mapping-library/{record_id}")
async def api_mapping_library_update(record_id: str, body: MappingLibItem, user: dict = Depends(auth.require_login)):
    """更新一条映射记录"""
    ok = update_mapping_record(record_id, body.source, body.target,
                               on_duplicate=body.on_duplicate)
    if not ok:
        raise HTTPException(status_code=404, detail="映射记录不存在或源列名冲突")
    return {"success": True}


@app.delete("/api/mapping-library/{record_id}")
async def api_mapping_library_delete(record_id: str, user: dict = Depends(auth.require_login)):
    """删除一条映射记录"""
    ok = delete_mapping_record(record_id)
    if not ok:
        raise HTTPException(status_code=404, detail="映射记录不存在")
    return {"success": True}


@app.delete("/api/mapping-library")
async def api_mapping_library_clear(user: dict = Depends(auth.require_login)):
    """清空整个映射库"""
    count = clear_mapping_records()
    return {"success": True, "deleted": count}


@app.get("/api/mapping-library/export")
def api_mapping_library_export(user: dict = Depends(auth.require_login)):
    """导出映射库全部记录（Excel 表格 .xlsx）"""
    import io
    import pandas as pd
    from datetime import date
    from urllib.parse import quote

    records = list_mapping_records()
    source_label = {"ai": "AI自动", "manual": "手动", "import": "导入"}
    rows = []
    for r in records:
        rows.append({
            "源列名": r.get("source_column", ""),
            "目标列名": r.get("target_column", ""),
            "来源": source_label.get(r.get("source", ""), r.get("source", "") or ""),
            "使用次数": r.get("use_count", 0),
            "最近使用": r.get("last_used", ""),
            "创建时间": r.get("created_at", ""),
        })

    df = pd.DataFrame(rows)
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)

    filename = f"映射匹配库_{date.today().isoformat()}.xlsx"
    encoded = quote(filename)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )


@app.post("/api/mapping-library/import/parse")
def api_mapping_library_import_parse(file: UploadFile = File(...), user: dict = Depends(auth.require_login)):
    """解析上传的表格文件（.xlsx/.xls/.csv），提取映射记录"""
    import io
    import pandas as pd

    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名为空")
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in (".xlsx", ".xls", ".csv"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xls / .csv 格式")

    content = file.file.read()
    try:
        if ext == ".csv":
            df = pd.read_csv(io.BytesIO(content), dtype=object)
        else:
            df = pd.read_excel(io.BytesIO(content), dtype=object)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败: {e}")

    if df is None or df.empty:
        raise HTTPException(status_code=400, detail="文件中没有数据")

    # 识别"源列名 / 目标列名"列（兼容常见命名）
    cols = [str(c).strip() if not pd.isna(c) else "" for c in df.columns]
    src_col_idx = None
    tgt_col_idx = None
    for i, c in enumerate(cols):
        cl = c.lower()
        if src_col_idx is None and ("源" in c or "source" in cl):
            src_col_idx = i
        if tgt_col_idx is None and ("目标" in c or "target" in cl or "参考" in cl):
            tgt_col_idx = i
    if src_col_idx is None and tgt_col_idx is None and len(cols) >= 2:
        src_col_idx = 0
        tgt_col_idx = 1

    records = []
    for _, row in df.iterrows():
        src = row.iloc[src_col_idx] if src_col_idx is not None else None
        tgt = row.iloc[tgt_col_idx] if tgt_col_idx is not None else None
        if src is None or tgt is None:
            continue
        src_s = str(src).strip() if not pd.isna(src) else ""
        tgt_s = str(tgt).strip() if not pd.isna(tgt) else ""
        if src_s and tgt_s:
            records.append({"source": src_s, "target": tgt_s})

    if not records:
        raise HTTPException(
            status_code=400,
            detail="未从文件中解析到有效的映射记录（表格需包含「源列名」「目标列名」列）",
        )
    return {"success": True, "records": records, "count": len(records)}


@app.post("/api/mapping-library/import")
async def api_mapping_library_import(body: ImportRequest, user: dict = Depends(auth.require_login)):
    """批量导入映射记录"""
    if not body.records:
        raise HTTPException(status_code=400, detail="导入内容为空")
    count = import_mapping_records(
        [{"source_column": r.source, "target_column": r.target} for r in body.records],
        mode=body.mode,
        on_duplicate=body.on_duplicate,
    )
    return {"success": True, "imported": count}


@app.post("/api/mapping-library/batch-delete")
async def api_mapping_library_batch_delete(body: BatchDeleteRequest, user: dict = Depends(auth.require_login)):
    """批量删除映射记录"""
    deleted = batch_delete_mapping_records(body.ids)
    return {"success": True, "deleted": deleted}


def _default_formula_action(is_src_formula: bool, is_tgt_formula: bool) -> str:
    """根据源/目标列是否为公式列，给出默认的公式处理方式"""
    if is_src_formula and is_tgt_formula:
        return "target_formula"  # 双方都是公式 → 默认保留参考公式
    if is_tgt_formula:
        return "target_formula"  # 参考表公式列 → 默认保留参考公式
    if is_src_formula:
        return "source_data"     # 源表公式列 → 默认写源数据（计算值）
    return "default"


@app.post("/api/match")
def api_match(request: Request, body: MatchRequest = MatchRequest(), user: dict = Depends(auth.require_login)):
    """执行 AI 列名匹配（多文件时并发调用，缩短总等待时间）

    同步 def 路由：内部会调用 DeepSeek API（可能长时间阻塞）。
    必须由线程池执行，绝不能阻塞 asyncio 事件循环，否则单个 AI 请求挂起
    会让整个服务（含登录）全部卡死（2026-08-07 线上事故根因）。
    """
    import pandas as pd
    from concurrent.futures import ThreadPoolExecutor, as_completed

    data = get_session(request)

    if not data["source_files"]:
        raise HTTPException(status_code=400, detail="请先上传原始表格")
    if not data["reference_path"]:
        raise HTTPException(status_code=400, detail="请先上传参考模板")

    use_ai = body.useAI

    # 若使用 AI 但未配置 API Key，直接提示用户先配置（前端也会强制引导）
    if use_ai and not ai_config.has_api_key():
        raise HTTPException(
            status_code=400,
            detail="尚未配置 DeepSeek API Key，请点击右上角「⚙ API 设置」填写后再使用 AI 匹配",
        )

    # 重置本次匹配的 token 累计（多次文件/多次调用累加）
    ai_config.reset_match_usage()

    try:
        template = read_reference_template(data["reference_path"])
        target_headers = template["headers"]
        data["reference_headers"] = target_headers
        formula_headers = get_formula_headers(template)

        all_mappings_list = []
        combined_mappings = []
        seen_mappings = set()

        ref_df = None
        try:
            ref_df = pd.read_excel(data["reference_path"], header=0, dtype=object)
            ref_df = ref_df.dropna(how="all").reset_index(drop=True)
        except Exception:
            pass

        def _match_single(idx, src_info):
            """匹配单个文件，返回 (index, mappings)"""
            df, source_headers, _ = read_source_file(src_info["path"])
            mappings = match_headers(
                source_headers=source_headers,
                target_headers=target_headers,
                use_ai=use_ai,
                source_df=df,
                reference_df=ref_df,
                reference_headers=target_headers,
                formula_columns=formula_headers,
            )
            return idx, mappings

        source_files = data["source_files"]

        # 多个源文件 → 并发匹配（每个文件的 AI 调用并行），显著减少总耗时
        results = {}
        if len(source_files) > 1:
            workers = min(3, len(source_files))  # 并发上限 3，避免 API 限流
            with ThreadPoolExecutor(max_workers=workers) as executor:
                futures = [
                    executor.submit(_match_single, i, src_info)
                    for i, src_info in enumerate(source_files)
                ]
                for fut in as_completed(futures):
                    idx, mappings = fut.result()
                    results[idx] = mappings
        else:
            idx, mappings = _match_single(0, source_files[0])
            results[0] = mappings

        # 按原始顺序处理结果，保证输出稳定
        target_formula_set = set(formula_headers)  # 参考表公式列名
        for idx, src_info in enumerate(source_files):
            mappings = results[idx]
            source_formula_set = set(src_info.get("source_formula_columns", []))  # 源表公式列名
            editable = []
            mapped_sources = set()
            for m in mappings:
                src_col = m["source_column"]
                tgt_col = m["target_column"]
                is_src_formula = normalize_header(src_col) in source_formula_set
                is_tgt_formula = normalize_header(tgt_col) in target_formula_set
                editable.append({
                    "source": src_col,
                    "target": tgt_col,
                    "confidence": m.get("confidence", "medium"),
                    "reason": m.get("reason", ""),
                    "source_is_formula": is_src_formula,
                    "target_is_formula": is_tgt_formula,
                    "formula_action": _default_formula_action(is_src_formula, is_tgt_formula),
                    "add_to_library": True,   # 匹配结果默认勾选「加入映射库」
                })
                mapped_sources.add(src_col)
                key = (src_col, tgt_col)
                if key not in seen_mappings:
                    seen_mappings.add(key)
                    combined_mappings.append(editable[-1])

            unmapped = [h for h in src_info["headers"] if h not in mapped_sources]
            all_mappings_list.append({
                "source_name": src_info["name"],
                "mappings": editable,
                "unmapped": unmapped,
            })

        data["mappings_list"] = all_mappings_list
        data["original_mappings"] = [
            {"source": m["source"], "target": m["target"]}
            for m in combined_mappings
        ]

        all_source_headers = []
        seen_headers = set()
        for src_info in data["source_files"]:
            for h in src_info["headers"]:
                if h not in seen_headers:
                    seen_headers.add(h)
                    all_source_headers.append(h)

        all_unmapped = []
        seen_unmapped = set()
        for item in all_mappings_list:
            for h in item["unmapped"]:
                if h not in seen_unmapped:
                    seen_unmapped.add(h)
                    all_unmapped.append(h)

        # 汇总源表公式列名（去重），供前端在用户修改/添加映射时动态标注公式列
        all_source_formula_cols = []
        seen_sf = set()
        for src_info in data["source_files"]:
            for fc in src_info.get("source_formula_columns", []):
                if fc not in seen_sf:
                    seen_sf.add(fc)
                    all_source_formula_cols.append(fc)
        all_target_formula_cols = []
        seen_tf = set()
        for fc in formula_headers:
            if fc not in seen_tf:
                seen_tf.add(fc)
                all_target_formula_cols.append(fc)

        return {
            "success": True,
            "sessionId": get_session_id(request),
            "sourceFiles": [{"name": f["name"], "rows": f["rows"]} for f in data["source_files"]],
            "mappings": combined_mappings,
            "unmappedSources": all_unmapped,
            "targetOptions": target_headers,
            "sourceHeaders": all_source_headers,
            "allMappingsList": all_mappings_list,
            "sourceFormulaColumns": all_source_formula_cols,
            "targetFormulaColumns": all_target_formula_cols,
            "aiUsage": {"totalTokens": ai_config.get_match_usage()},
        }
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"匹配失败: {str(e)}")


@app.post("/api/convert")
def api_convert(request: Request, body: ConvertRequest = ConvertRequest(), user: dict = Depends(auth.require_login)):
    """执行表格转换（支持多源文件批量处理并合并）

    同步 def 路由：openpyxl/pandas 转换大文件可能耗时较长，
    放入线程池执行以免阻塞事件循环。
    """
    data = get_session(request)

    if not data["source_files"]:
        raise HTTPException(status_code=400, detail="请先上传原始表格")
    if not data["reference_path"]:
        raise HTTPException(status_code=400, detail="请先上传参考模板")

    user_mappings = body.mappings

    try:
        source_files = data["source_files"]

        image_diag = {}
        if len(source_files) == 1:
            src = source_files[0]
            output_filename = f"{src['name']}_转换结果.xlsx"
            output_path = os.path.join(OUTPUT_DIR, output_filename)

            converted_mappings = []
            if user_mappings:
                converted_mappings = [{
                    "source_column": m.source if m.source else "",
                    "target_column": m.target if m.target else "",
                    "formula_action": m.formula_action,
                } for m in user_mappings if m.target]

            result_path = _do_convert(
                source_path=src["path"],
                reference_path=data["reference_path"],
                output_path=output_path,
                mappings=converted_mappings,
                source_name=src["name"],
                source_formula_info=src.get("source_formula_info", {}),
                image_diag=image_diag,
            )
        else:
            output_filename = f"批量转换结果_{len(source_files)}个文件.xlsx"
            output_path = os.path.join(OUTPUT_DIR, output_filename)
            image_diags = [{} for _ in source_files]
            result_path = _do_batch_convert(
                source_files=source_files,
                reference_path=data["reference_path"],
                output_path=output_path,
                user_mappings=user_mappings,
                image_diags=image_diags,
            )
            # 批量：任一文件图片只能按顺序推断 → 整体提示
            if any(d.get('mode') == 'fallback' for d in image_diags):
                image_diag = {
                    'mode': 'fallback',
                    'count': sum(d.get('count', 0) for d in image_diags),
                }

        data["output_path"] = result_path
        data["output_name"] = os.path.basename(result_path)

        # 保存映射到匹配库：
        #  1. 未被修改的 AI 匹配结果 → 自动入库
        #  2. 用户勾选了「加入映射库」的映射 → 入库
        #  重复处理方式由前端在转换前确认后传入（overwrite/skip/abort）
        if user_mappings:
            original_mappings = data.get("original_mappings", [])
            try:
                stats = save_convert_mappings(
                    mappings=[
                        {
                            "source": m.source,
                            "target": m.target,
                            "add_to_library": m.add_to_library,
                        }
                        for m in user_mappings
                    ],
                    original_mappings=original_mappings,
                    on_duplicate=body.on_duplicate,
                )
                print(f"  📚 映射入库统计: {stats}")
            except Exception as e:
                print(f"  ⚠️ 保存映射到匹配库失败: {e}")

        first_source = source_files[0]["path"]
        comparison = _build_comparison(first_source, result_path)

        return {
            "success": True,
            "sessionId": get_session_id(request),
            "outputName": os.path.basename(result_path),
            "downloadUrl": "/api/download",
            "comparison": comparison,
            "imageWarning": _build_image_warning(source_files, image_diag),
        }
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"转换失败: {str(e)}")


# ============================================================
# AI 配置（API Key 加密管理 + token 消耗统计）
# ============================================================

@app.get("/api/ai/config")
async def api_ai_config_get(request: Request, user: dict = Depends(auth.require_login)):
    """读取当前 AI 配置状态（是否已配置 Key、掩码、当前 Key 总消耗等）"""
    data = get_session(request)
    if not data.get("ai_nonce"):
        data["ai_nonce"] = secrets.token_hex(16)
    return {
        "success": True,
        "hasKey": ai_config.has_api_key(),
        "maskedKey": ai_config.masked_api_key(),
        "totalTokens": ai_config.get_total_tokens(),
        "model": DEEPSEEK_MODEL,
        "url": DEEPSEEK_API_URL,
        "nonce": data["ai_nonce"],
    }


@app.post("/api/ai/config")
async def api_ai_config_save(request: Request, body: AIConfigRequest | None = None, user: dict = Depends(auth.require_login)):
    """保存（加密）新的 API Key，并重置该 Key 的消耗统计"""
    data = get_session(request)
    nonce = data.get("ai_nonce", "")
    api_key = ""
    if body and body.api_key_enc:
        try:
            api_key = _transport_decrypt(body.api_key_enc, nonce)
        except Exception:
            raise HTTPException(status_code=400, detail="加密数据无效，请刷新页面后重试")
    elif body and body.api_key:
        api_key = body.api_key
    api_key = (api_key or "").strip()
    if not api_key:
        raise HTTPException(status_code=400, detail="API Key 不能为空")
    if not api_key.startswith("sk-"):
        raise HTTPException(status_code=400, detail="API Key 格式不正确（应以 sk- 开头）")
    ai_config.save_api_key(api_key)
    # 保存后重新生成 nonce，防止旧 nonce 复用
    data["ai_nonce"] = secrets.token_hex(16)
    return {
        "success": True,
        "hasKey": True,
        "maskedKey": ai_config.masked_api_key(),
        "totalTokens": ai_config.get_total_tokens(),
    }


@app.get("/api/download")
async def download_file(request: Request, user: dict = Depends(auth.require_login)):
    """下载转换后的文件"""
    data = get_session(request)
    filepath = data.get("output_path", "")

    if not filepath or not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="文件不存在，请先执行转换")

    download_name = data.get("output_name", "result.xlsx")
    return FileResponse(
        path=filepath,
        filename=download_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/reset")
async def api_reset(request: Request, user: dict = Depends(auth.require_login)):
    """重置会话"""
    sid = get_session_id(request)
    if sid in sessions:
        sess = sessions[sid]
        for key in ("source_paths", "source_files"):
            items = sess.get(key, [])
            if isinstance(items, list):
                for item in items:
                    p = item if isinstance(item, str) else item.get("path", "")
                    if p and os.path.exists(p):
                        try:
                            os.remove(p)
                        except Exception:
                            pass
        if sess.get("reference_path") and os.path.exists(sess["reference_path"]):
            try:
                os.remove(sess["reference_path"])
            except Exception:
                pass
        del sessions[sid]

    return {"success": True}


# ============================================================
# 辅助函数（从原 Flask app.py 迁移）
# ============================================================

def _build_comparison(source_path, output_path):
    try:
        import openpyxl
        from table_converter import dedupe_headers
        src_df, src_headers, _ = read_source_file(source_path)
        src_preview = []
        max_rows = min(20, len(src_df))
        for i in range(max_rows):
            row = {}
            for h in src_headers:
                val = src_df.iloc[i][h] if h in src_df.columns else ""
                if hasattr(val, "item"):
                    val = val.item()
                if val is not None and str(val) != "nan":
                    row[str(h)] = str(val)
                else:
                    row[str(h)] = ""
            src_preview.append(row)

        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
        out_headers = []
        for c in range(1, ws.max_column + 1):
            h = ws.cell(row=1, column=c).value
            out_headers.append(str(h) if h else "")
        # 输出文件表头保留原始列名；预览时编号去重，避免重复列名 key 冲突
        out_headers_display = dedupe_headers(out_headers)

        out_preview = []
        out_max_rows = min(20, ws.max_row - 1)
        for r in range(2, out_max_rows + 2):
            row = {}
            for c in range(1, ws.max_column + 1):
                h = out_headers_display[c - 1] if c - 1 < len(out_headers_display) else f"Col{c}"
                val = ws.cell(row=r, column=c).value
                row[str(h)] = str(val) if val is not None else ""
            out_preview.append(row)
        wb.close()

        return {
            "sourceHeaders": src_headers,
            "sourcePreview": src_preview,
            "sourceRowCount": len(src_df),
            "outputHeaders": out_headers_display,
            "outputPreview": out_preview,
            "outputRowCount": ws.max_row - 1,
        }
    except Exception as e:
        return {"error": str(e)}


# ============================================================
# 图片列识别与嵌入放置（图片转换健壮性）
# ============================================================
def is_image_column_name(header: str) -> bool:
    """判断列名是否为图片列（兼容中英文常见命名，如 图片/照片/图1/主图/Image/Img/Photo）"""
    h = normalize_header(header).lower()
    if not h:
        return False
    for kw in ("图片", "照片", "image", "photo", "img"):
        if kw in h:
            return True
    if h == "图":
        return True
    if h.startswith("图") and len(h) > 1 and h[1].isdigit():
        return True
    return False


def find_source_image_columns(images_map: dict, source_headers: list) -> list:
    """从已提取的图片位置推断源表图片列。

    优先返回「表头名称为图片列且实际有图片」的列（按图片数量降序），
    其次返回其他实际含有图片的列（同样按数量降序），供调用方取第一个作为主图片列。
    这样即使源表图片列名不含"图片"（如 主图/照片/IMG），也能稳定定位图片列。
    """
    counts = {}
    for (_row, col) in images_map:
        counts[col] = counts.get(col, 0) + 1
    if not counts:
        return []
    named, unnamed = [], []
    for col, cnt in counts.items():
        h = normalize_header(source_headers[col]) if col < len(source_headers) else ""
        (named if is_image_column_name(h) else unnamed).append((col, cnt))
    named.sort(key=lambda x: -x[1])
    unnamed.sort(key=lambda x: -x[1])
    return [c for c, _ in named] + [c for c, _ in unnamed]


def _embed_images(ws, images_map, source_img_col, target_img_col, data_rows,
                  col_width_emu, row_height_pt, margin_pt, row_offset=0):
    """将源表图片以「嵌入单元格」方式放置到目标工作表图片列。

    - 源表无论浮动图(drawings/锚点)还是嵌入图(cellImages/DISPIMG)，均已由
      extract_images_from_source 统一提取为字节数据。
    - 输出使用 TwoCellAnchor：图片严格落在目标单元格内、保持纵横比、居中放置，
      且随单元格移动/缩放，视觉与行为上都表现为嵌入单元格（而非漂浮在表上）。
    - row_offset: 批量合并时本文件首数据行相对输出第 2 行的偏移行数（单文件为 0）。
    """
    import io as _io
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor

    row_h_emu = int(row_height_pt * 12700)
    max_h_emu = int((row_height_pt - margin_pt * 2) * 12700)
    if col_width_emu <= 0:
        col_width_emu = 1143000

    for (src_row, src_col), img_data in images_map.items():
        if src_col != source_img_col or src_row >= data_rows:
            continue
        try:
            probe = XLImage(_io.BytesIO(img_data))
            ow = max(probe.width, 1)
            oh = max(probe.height, 1)
            ow_emu = int(ow * 914400 / 96)
            oh_emu = int(oh * 914400 / 96)
            # 先按行高上限缩放
            if oh_emu > max_h_emu:
                ratio = max_h_emu / oh_emu
                sh_emu = max_h_emu
                sw_emu = int(ow_emu * ratio)
            else:
                sh_emu = oh_emu
                sw_emu = ow_emu
            # 再按列宽上限缩放，避免图片超出单元格（嵌入效果更稳定）
            if sw_emu > col_width_emu:
                ratio = col_width_emu / sw_emu
                sw_emu = col_width_emu
                sh_emu = int(sh_emu * ratio)
        except Exception:
            continue
        if sw_emu <= 0 or sh_emu <= 0:
            continue

        col_off = (col_width_emu - sw_emu) // 2
        row_off = (row_h_emu - sh_emu) // 2
        anchor_row = src_row + 1 + row_offset   # 0-based 目标数据行（输出数据行 = src_row+2+row_offset）
        try:
            img = XLImage(_io.BytesIO(img_data))
            from_marker = AnchorMarker(col=target_img_col, colOff=col_off, row=anchor_row, rowOff=row_off)
            to_marker = AnchorMarker(col=target_img_col, colOff=col_off + sw_emu, row=anchor_row, rowOff=row_off + sh_emu)
            ws.add_image(img, TwoCellAnchor(_from=from_marker, to=to_marker))
        except Exception:
            continue


def _build_image_warning(source_files: list, image_diag: dict):
    """当 .xls 源文件的图片只能按存储顺序推断位置时，生成前端提醒；否则返回 None。

    image_diag 由 extract_images_from_source 写入：
      mode: exact（精确位置）/ fallback（顺序推断兜底）/ none（无图片）
    """
    if not image_diag or image_diag.get('mode') != 'fallback':
        return None
    if not image_diag.get('count'):
        return None
    xls_files = [s['name'] for s in source_files
                 if str(s.get('name', '')).lower().endswith('.xls')]
    if not xls_files:
        return None
    return {
        "files": xls_files,
        "imageCount": image_diag.get('count', 0),
        "message": (
            "检测到旧版 .xls 格式文件中的图片缺少单元格锚点信息，"
            "转换结果里的图片只能按存储顺序逐行放置，可能与商品行错位。"
            "建议在 WPS/Excel 中将源文件另存为 .xlsx 后重新上传，图片位置即可精确还原。"
        ),
    }


def _do_convert(source_path, reference_path, output_path, mappings, source_name,
                source_formula_info=None, image_diag=None):
    import pandas as pd
    import openpyxl
    from openpyxl.utils import get_column_letter
    from copy import copy
    import re

    if source_formula_info is None:
        source_formula_info = {}

    df, source_headers, source_header_row = read_source_file(source_path)
    template = read_reference_template(reference_path)
    target_headers = template["headers"]
    formula_info = template["formula_info"]
    num_target_cols = template["max_column"]

    column_mapping = {}          # 源列名 → 目标列索引
    formula_actions = {}         # 目标列索引 → 公式处理方式
    source_src_formula = {}      # 源列名 → 源表公式（当选择写源公式时）

    same_format = (
        len(source_headers) == len(target_headers)
        and sum(
            1 for sh in source_headers
            if any(normalize_header(sh) == normalize_header(th) for th in target_headers)
        ) >= len(source_headers) * 0.7
    )

    def _resolve_mapping(m):
        """解析单个映射，返回 (src_col, tgt_idx, action) 或 None"""
        src_col = m.get("source_column", "")
        tgt_col = m.get("target_column", "")
        action = m.get("formula_action", "default")
        if not src_col or not tgt_col:
            return None
        src_idx = None
        for i, h in enumerate(source_headers):
            if normalize_header(h) == normalize_header(src_col):
                src_idx = i
                break
        if src_idx is None:
            return None
        tgt_idx = None
        tgt_clean = normalize_header(tgt_col)
        for idx, h in enumerate(target_headers):
            if normalize_header(h) == tgt_clean:
                tgt_idx = idx
                break
        if tgt_idx is None:
            for idx, h in enumerate(target_headers):
                h_clean = normalize_header(h)
                if tgt_clean in h_clean or h_clean in tgt_clean:
                    tgt_idx = idx
                    break
        if tgt_idx is None:
            return None

        # 源表公式列：确定源列字母对应的公式
        src_formula = None
        if action == "source_formula":
            src_col_letter = get_column_letter(src_idx + 1)
            src_formula = source_formula_info.get(src_col_letter, "")

        return {
            "src_col": src_col,
            "src_idx": src_idx,
            "tgt_idx": tgt_idx,
            "action": action,
            "src_formula": src_formula,
        }

    if same_format:
        # 即使两表格式相同，也优先应用用户指定的映射与公式动作
        for m in mappings:
            r = _resolve_mapping(m)
            if r is None:
                continue
            column_mapping[r["src_col"]] = r["tgt_idx"]
            formula_actions[r["tgt_idx"]] = r["action"]
            if r["src_formula"]:
                source_src_formula[r["src_col"]] = r["src_formula"]
        # 补充：未在映射中指定的列，按位置一一对应
        for i, sh in enumerate(source_headers):
            if i < len(target_headers) and i not in formula_actions:
                column_mapping[sh] = i
                formula_actions[i] = "default"
    else:
        for m in mappings:
            r = _resolve_mapping(m)
            if r is None:
                continue
            column_mapping[r["src_col"]] = r["tgt_idx"]
            formula_actions[r["tgt_idx"]] = r["action"]
            if r["src_formula"]:
                source_src_formula[r["src_col"]] = r["src_formula"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # 输出文件表头使用原始列名（重复列名不编号，编号仅在前端区分）
    output_headers = template.get("raw_headers") or template["headers"]
    for col_idx in range(1, num_target_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = output_headers[col_idx - 1] if col_idx - 1 < len(output_headers) else target_headers[col_idx - 1]
        col_letter = get_column_letter(col_idx)
        if col_letter in template["header_styles"]:
            styles = template["header_styles"][col_letter]
            if styles.get("font"):
                cell.font = styles["font"]
            if styles.get("alignment"):
                cell.alignment = styles["alignment"]
            if styles.get("fill"):
                cell.fill = styles["fill"]
            if styles.get("border"):
                cell.border = styles["border"]

    def _translate_source_formula(formula, output_row):
        """将源表公式改写为输出表适用的公式：
        - 源列字母 → 对应目标列字母（按列映射关系）
        - 数据行号 → output_row（每行数据独立）
        """
        src_idx_of = {}
        for i, h in enumerate(source_headers):
            src_idx_of[normalize_header(h)] = i
        letter_map = {}
        for src_name, tgt_idx in column_mapping.items():
            src_idx = src_idx_of.get(normalize_header(src_name))
            if src_idx is None:
                continue
            letter_map[get_column_letter(src_idx + 1)] = get_column_letter(tgt_idx + 1)

        def repl(match):
            cl = match.group(1)
            rn = int(match.group(2))
            new_cl = letter_map.get(cl, cl)
            return f"{new_cl}{output_row}" if rn >= 2 else f"{new_cl}{rn}"

        return re.sub(r"([A-Z]+)(\d+)", repl, formula)

    def _adjust_ref_formula(formula, output_row):
        """参考公式行号改写：数据行(2) → output_row"""
        def replace_row(match):
            cl = match.group(1)
            rn = int(match.group(2))
            return f"{cl}{output_row}" if rn == 2 else match.group(0)
        return re.sub(r"([A-Z]+)(\d+)", replace_row, formula)

    for src_row_idx in range(len(df)):
        output_row = src_row_idx + 2

        # --- 第一遍：写入源数据映射的值（跳过要写公式的目标列） ---
        for col_idx in range(1, num_target_cols + 1):
            cell = ws.cell(row=output_row, column=col_idx)
            col_letter = get_column_letter(col_idx)
            if col_letter in template["data_styles"]:
                styles = template["data_styles"][col_letter]
                if styles.get("font"):
                    cell.font = styles["font"]
                if styles.get("alignment"):
                    cell.alignment = styles["alignment"]
                if styles.get("fill"):
                    cell.fill = styles["fill"]
                if styles.get("border"):
                    cell.border = styles["border"]
                if styles.get("number_format"):
                    cell.number_format = styles["number_format"]

            action = formula_actions.get(col_idx - 1, "default")
            # 目标列要写公式（保留参考公式 / 写源公式）时，第一遍不写数据
            if action in ("target_formula", "source_formula"):
                continue

            src_col_idx = None
            for src_col_name, tgt_idx in column_mapping.items():
                if tgt_idx == col_idx - 1:
                    for i, h in enumerate(source_headers):
                        if normalize_header(h) == normalize_header(src_col_name):
                            src_col_idx = i
                            break
                    break

            if src_col_idx is not None and src_col_idx < len(df.columns):
                val = df.iloc[src_row_idx, src_col_idx]
                if pd.notna(val):
                    val_str = str(val)
                    if val_str.startswith("=DISPIMG"):
                        cell.value = None
                    else:
                        try:
                            if isinstance(val, str) and val_str.strip() == "":
                                cell.value = None
                            elif isinstance(val, (int, float)):
                                if isinstance(val, float):
                                    rounded = round(val, 6)
                                    cell.value = int(rounded) if rounded == int(rounded) else rounded
                                else:
                                    cell.value = val
                            else:
                                cleaned = val_str.replace(",", "").strip()
                                if cleaned.replace(".", "").replace("-", "").isdigit():
                                    num_val = float(cleaned)
                                    rounded = round(num_val, 6)
                                    cell.value = int(rounded) if rounded == int(rounded) else rounded
                                else:
                                    cell.value = val
                        except (ValueError, TypeError):
                            cell.value = val

        # --- 第二遍：按公式处理方式写公式 ---
        for col_idx in range(1, num_target_cols + 1):
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=output_row, column=col_idx)
            action = formula_actions.get(col_idx - 1, "default")

            # 找到映射到此目标列的源列（用于取源公式）
            src_col_name_for_tgt = None
            for scn, tidx in column_mapping.items():
                if tidx == col_idx - 1:
                    src_col_name_for_tgt = scn
                    break

            if action == "source_formula" and src_col_name_for_tgt:
                # 用户选择：写源公式
                src_formula = source_src_formula.get(src_col_name_for_tgt, "")
                if src_formula:
                    cell.value = _translate_source_formula(src_formula, output_row)
                    continue

            if col_letter in formula_info:
                # 参考公式列：
                #   target_data → 不写公式（第一遍已写源数据）
                #   其他（default / target_formula / 无映射）→ 保留参考公式
                if action == "target_data":
                    continue
                formula = formula_info[col_letter]
                cell.value = _adjust_ref_formula(formula, output_row)

    images_map = extract_images_from_source(
        source_path, header_row=source_header_row, diag=image_diag,
    )

    # ---- 稳健识别源/目标图片列（不再只依赖表头含"图片"，优先从实际图片位置推断） ----
    source_img_col = None
    _img_cands = find_source_image_columns(images_map, source_headers)
    if _img_cands:
        source_img_col = _img_cands[0]
    if source_img_col is None:
        for i, h in enumerate(source_headers):
            if is_image_column_name(h):
                source_img_col = i
                break

    target_img_col = None
    if source_img_col is not None:
        _src_img_name = normalize_header(source_headers[source_img_col])
        for scn, tgt_idx in column_mapping.items():
            if normalize_header(scn) == _src_img_name:
                target_img_col = tgt_idx
                break
    if target_img_col is None:
        for i, h in enumerate(target_headers):
            if is_image_column_name(h):
                target_img_col = i
                break

    if images_map and source_img_col is not None and target_img_col is not None:
        target_col_letter = get_column_letter(target_img_col + 1)
        col_width_chars = template["column_widths"].get(target_col_letter, None)
        col_width_emu = int(col_width_chars * 7 * 914400 / 96) if col_width_chars else 1143000
        for r in range(2, len(df) + 2):
            ws.row_dimensions[r].height = IMAGE_ROW_HEIGHT_PT
        _embed_images(ws, images_map, source_img_col, target_img_col, len(df),
                      col_width_emu, IMAGE_ROW_HEIGHT_PT, IMAGE_MARGIN_PT)

    for col_idx in range(1, num_target_cols + 1):
        col_letter = get_column_letter(col_idx)
        if col_letter in template["column_widths"] and template["column_widths"][col_letter]:
            ws.column_dimensions[col_letter].width = template["column_widths"][col_letter]

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)   # 输出目录被删除后自动重建
    wb.save(output_path)
    wb.close()
    return output_path


def _do_batch_convert(source_files, reference_path, output_path, user_mappings=None,
                      image_diags=None):
    import pandas as pd
    import openpyxl
    from openpyxl.utils import get_column_letter
    from copy import copy
    import re

    template = read_reference_template(reference_path)
    target_headers = template["headers"]
    num_target_cols = template["max_column"]
    target_img_col = None
    for i, h in enumerate(target_headers):
        if is_image_column_name(h):
            target_img_col = i
            break

    if target_img_col is not None:
        tgt_letter = get_column_letter(target_img_col + 1)
        col_width_chars = template["column_widths"].get(tgt_letter, None)
        col_width_emu = int(col_width_chars * 7 * 914400 / 96) if col_width_chars else 1143000
    else:
        col_width_emu = 1143000

    ROW_HEIGHT_PT = IMAGE_ROW_HEIGHT_PT
    MARGIN_PT = IMAGE_MARGIN_PT
    MAX_IMG_H_EMU = int((ROW_HEIGHT_PT - MARGIN_PT * 2) * 12700)
    ROW_HEIGHT_EMU = int(ROW_HEIGHT_PT * 12700)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # 输出文件表头使用原始列名（重复列名不编号，编号仅在前端区分）
    output_headers = template.get("raw_headers") or template["headers"]
    for col_idx in range(1, num_target_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = output_headers[col_idx - 1] if col_idx - 1 < len(output_headers) else target_headers[col_idx - 1]
        col_letter = get_column_letter(col_idx)
        if col_letter in template["header_styles"]:
            styles = template["header_styles"][col_letter]
            if styles.get("font"): cell.font = styles["font"]
            if styles.get("alignment"): cell.alignment = styles["alignment"]
            if styles.get("fill"): cell.fill = styles["fill"]
            if styles.get("border"): cell.border = styles["border"]

    current_output_row = 2

    for file_idx, src_info in enumerate(source_files):
        source_path = src_info["path"]
        source_name = src_info["name"]
        temp_output = os.path.join(OUTPUT_DIR, f"_temp_batch_{file_idx}.xlsx")

        if user_mappings:
            mappings = [{
                "source_column": m.source if hasattr(m, "source") else m.get("source", ""),
                "target_column": m.target if hasattr(m, "target") else m.get("target", ""),
                "formula_action": m.formula_action if hasattr(m, "formula_action") else "default",
            } for m in user_mappings if (m.target if hasattr(m, "target") else m.get("target", ""))]
        else:
            df, src_headers, _ = read_source_file(source_path)
            ref_df = None
            try:
                ref_df = pd.read_excel(reference_path, header=0, dtype=object)
                ref_df = ref_df.dropna(how="all").reset_index(drop=True)
            except Exception:
                pass
            mappings_raw = match_headers(
                src_headers, target_headers, use_ai=True,
                source_df=df, reference_df=ref_df,
                reference_headers=target_headers,
                formula_columns=get_formula_headers(template),
            )
            mappings = [{
                "source_column": m["source_column"],
                "target_column": m["target_column"],
                "formula_action": "default",
            } for m in mappings_raw]

        _do_convert(source_path, reference_path, temp_output, mappings, source_name,
                    source_formula_info=src_info.get("source_formula_info", {}))

        temp_wb = openpyxl.load_workbook(temp_output)
        temp_ws = temp_wb.active
        rows_in_this_file = temp_ws.max_row - 1

        for row_idx in range(2, temp_ws.max_row + 1):
            dst_row = current_output_row
            for col_idx in range(1, num_target_cols + 1):
                src_cell = temp_ws.cell(row=row_idx, column=col_idx)
                dst_cell = ws.cell(row=dst_row, column=col_idx)
                val = src_cell.value
                if val is not None:
                    # 公式单元格：多文件合并时，后续文件的公式仍引用临时表行号（2..N），
                    # 需改写为合并后目标表行号，否则公式会引用错误的行
                    if isinstance(val, str) and val.startswith("="):
                        val = re.sub(
                            r"([A-Z]+)(\d+)",
                            lambda m: f"{m.group(1)}{dst_row}"
                                      if int(m.group(2)) == row_idx else m.group(0),
                            val,
                        )
                    dst_cell.value = val
                if src_cell.font: dst_cell.font = copy(src_cell.font)
                if src_cell.alignment: dst_cell.alignment = copy(src_cell.alignment)
                if src_cell.fill and src_cell.fill.start_color and src_cell.fill.start_color.rgb != "00000000":
                    dst_cell.fill = copy(src_cell.fill)
                if src_cell.border: dst_cell.border = copy(src_cell.border)
                if src_cell.number_format and src_cell.number_format != "General":
                    dst_cell.number_format = src_cell.number_format
            if row_idx in temp_ws.row_dimensions:
                h = temp_ws.row_dimensions[row_idx].height
                if h:
                    ws.row_dimensions[current_output_row].height = h
            current_output_row += 1

        df_src, src_headers, src_hr = read_source_file(source_path)
        images_map = extract_images_from_source(
            source_path,
            header_row=src_hr,
            diag=(image_diags[file_idx] if image_diags and file_idx < len(image_diags) else None),
        )

        # 稳健识别本文件的源图片列（优先从实际图片位置推断）
        source_img_col = None
        _cands = find_source_image_columns(images_map, src_headers)
        if _cands:
            source_img_col = _cands[0]
        if source_img_col is None:
            for i, h in enumerate(src_headers):
                if is_image_column_name(h):
                    source_img_col = i
                    break

        if target_img_col is not None and source_img_col is not None and images_map:
            row_offset = current_output_row - rows_in_this_file - 2
            _embed_images(ws, images_map, source_img_col, target_img_col, len(df_src),
                          col_width_emu, ROW_HEIGHT_PT, MARGIN_PT, row_offset=row_offset)

        temp_wb.close()
        try:
            os.remove(temp_output)
        except Exception:
            pass

    for col_idx in range(1, num_target_cols + 1):
        col_letter = get_column_letter(col_idx)
        if col_letter in template["column_widths"] and template["column_widths"][col_letter]:
            ws.column_dimensions[col_letter].width = template["column_widths"][col_letter]

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)   # 输出目录被删除后自动重建
    wb.save(output_path)
    wb.close()
    return output_path


# ============================================================
# 启动入口
# ============================================================
if __name__ == "__main__":
    import uvicorn
    print("\n" + "=" * 60)
    print("  📊 表格转换工具 - FastAPI 服务（React + FastAPI 架构）")
    print("=" * 60)
    print(f"  🌐 API 地址: http://127.0.0.1:8004")
    print(f"  📂 上传目录: {UPLOAD_DIR}")
    print(f"  📁 输出目录: {OUTPUT_DIR}")
    print(f"  🔓 CORS: 已启用")
    print(f"  📖 API 文档: http://127.0.0.1:8004/docs")
    print("=" * 60 + "\n")
    cleanup_old_files()
    uvicorn.run(app, host="0.0.0.0", port=8004)
