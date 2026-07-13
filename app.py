"""
表格转换工具 - Web 前端后端服务
基于 Flask 提供 REST API 和前端页面
"""

import os
import sys
import time
import uuid
from werkzeug.utils import secure_filename

from flask import (
    Flask, render_template, request, jsonify, 
    send_file, session
)

# 添加当前目录到路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from config import (
    BASE_DIR, OUTPUT_DIR,
    MARGIN_RATE,
    IMAGE_ROW_HEIGHT_PT,
    IMAGE_MARGIN_PT,
)
from table_converter import (
    read_source_file, read_reference_template,
    normalize_header, convert_table,
    extract_images_from_source, get_formula_headers,
)
from deepseek_matcher import match_headers
from mapping_history import batch_add_mappings, get_history_mappings_dict

# ============================================================
# Flask 应用初始化
# ============================================================
app = Flask(__name__)
app.secret_key = os.urandom(24).hex()
app.config['MAX_CONTENT_LENGTH'] = 300 * 1024 * 1024  # 300MB 上传限制（表格含图片较大）

# 上传和输出目录
UPLOAD_DIR = os.path.join(BASE_DIR, 'uploads')
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

# 文件清理配置：超过此时间的文件将被自动删除（秒）
CLEANUP_MAX_AGE = 24 * 3600  # 24 小时


def cleanup_old_files():
    """清理 uploads/ 和 输出/ 目录中过期的临时文件"""
    now = time.time()
    cleaned = 0
    
    for directory in (UPLOAD_DIR, OUTPUT_DIR):
        if not os.path.isdir(directory):
            continue
        for fname in os.listdir(directory):
            fpath = os.path.join(directory, fname)
            try:
                if os.path.isfile(fpath) and now - os.path.getmtime(fpath) > CLEANUP_MAX_AGE:
                    os.remove(fpath)
                    cleaned += 1
            except Exception:
                pass
    
    if cleaned:
        print(f"  🧹 已清理 {cleaned} 个过期临时文件")


# 会话存储（简化版，生产环境应用 Redis）
sessions = {}


def get_session():
    """获取或创建会话"""
    sid = session.get('sid')
    if not sid or sid not in sessions:
        sid = uuid.uuid4().hex
        session['sid'] = sid
        sessions[sid] = {
            'source_paths': [],       # 多文件路径列表
            'source_files': [],       # [{name, headers, rows, header_row, path, preview}]
            'reference_path': None,
            'reference_headers': [],
            'mappings_list': [],      # 每个源文件的映射 [{source_name, mappings: [...]}]
            'output_path': None,
        }
    return sessions[sid]


# ============================================================
# 错误处理 - 返回 JSON 而非 HTML
# ============================================================
@app.errorhandler(413)
def too_large(e):
    return jsonify({
        'success': False,
        'error': '文件过大，超过服务器上传限制(300MB)。请压缩表格中的图片后重试。'
    }), 413


@app.errorhandler(500)
def server_error(e):
    return jsonify({
        'success': False,
        'error': f'服务器内部错误: {str(e)}'
    }), 500


# ============================================================
# 页面路由
# ============================================================
@app.route('/')
def index():
    """主页面"""
    return render_template('index.html')


# ============================================================
# API 路由 - 文件上传
# ============================================================
@app.route('/api/upload/source', methods=['POST'])
def upload_source():
    """上传原始表格文件（支持多文件）"""
    cleanup_old_files()
    files = request.files.getlist('files')
    if not files or all(f.filename == '' for f in files):
        # 兼容旧版单文件上传
        if 'file' in request.files:
            files = [request.files['file']]
        else:
            return jsonify({'success': False, 'error': '未选择文件'})
    
    data = get_session()
    uploaded = []
    errors = []
    
    for file in files:
        if file.filename == '':
            continue
        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in ('.xlsx', '.xls'):
            errors.append(f'{file.filename}: 仅支持 .xlsx / .xls 格式')
            continue
        
        filename = secure_filename(f"source_{uuid.uuid4().hex[:8]}{ext}")
        filepath = os.path.join(UPLOAD_DIR, filename)
        file.save(filepath)
        
        try:
            df, headers, source_header_row = read_source_file(filepath)
            preview = df.head(3).fillna('').to_dict(orient='records')
            for row in preview:
                for k, v in row.items():
                    if isinstance(v, float):
                        row[k] = round(v, 4)
            
            file_info = {
                'name': os.path.splitext(file.filename)[0],
                'filename': file.filename,
                'path': filepath,
                'headers': headers,
                'rows': len(df),
                'header_row': source_header_row,
                'preview': preview,
            }
            data['source_files'].append(file_info)
            data['source_paths'].append(filepath)
            uploaded.append(file_info)
        except Exception as e:
            errors.append(f'{file.filename}: 解析失败 - {str(e)}')
    
    if not uploaded and errors:
        return jsonify({'success': False, 'error': '; '.join(errors)})
    
    return jsonify({
        'success': True,
        'files': [{'name': f['name'], 'filename': f['filename'],
                    'rows': f['rows'], 'headers': f['headers']} for f in data['source_files']],
        'totalFiles': len(data['source_files']),
        'newFiles': [{'name': f['name'], 'filename': f['filename'],
                       'rows': f['rows'], 'headers': f['headers']} for f in uploaded],
        'errors': errors if errors else None,
    })


@app.route('/api/upload/reference', methods=['POST'])
def upload_reference():
    """上传参考模板文件"""
    cleanup_old_files()
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': '未选择文件'})
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'error': '文件名为空'})
    
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ('.xlsx', '.xls'):
        return jsonify({'success': False, 'error': '仅支持 .xlsx / .xls 格式'})
    
    filename = secure_filename(f"ref_{uuid.uuid4().hex[:8]}{ext}")
    filepath = os.path.join(UPLOAD_DIR, filename)
    file.save(filepath)
    
    try:
        template = read_reference_template(filepath)
        
        data = get_session()
        data['reference_path'] = filepath
        data['reference_headers'] = template['headers']
        data['formula_count'] = len(template['formula_info'])
        data['fixed_value_count'] = len(template['fixed_values'])
        data['formula_info'] = {
            k: v for k, v in template['formula_info'].items()
        }
        
        return jsonify({
            'success': True,
            'filename': file.filename,
            'headers': template['headers'],
            'formulaColumns': list(template['formula_info'].keys()),
            'formulaCount': len(template['formula_info']),
        })
    except Exception as e:
        return jsonify({'success': False, 'error': f'模板解析失败: {str(e)}'})


# ============================================================
# API 路由 - 清空已上传文件
# ============================================================
@app.route('/api/clear/source', methods=['POST'])
def clear_source():
    """清空已上传的原始表格"""
    data = get_session()
    for src in data.get('source_files', []):
        try: os.remove(src['path'])
        except: pass
    for p in data.get('source_paths', []):
        try: os.remove(p)
        except: pass
    data['source_files'] = []
    data['source_paths'] = []
    data['mappings_list'] = []
    return jsonify({'success': True})


@app.route('/api/clear/reference', methods=['POST'])
def clear_reference():
    """清空已上传的参考模板"""
    data = get_session()
    if data.get('reference_path') and os.path.exists(data['reference_path']):
        try: os.remove(data['reference_path'])
        except: pass
    data['reference_path'] = None
    data['reference_headers'] = []
    data['mappings_list'] = []
    return jsonify({'success': True})


# ============================================================
# API 路由 - AI 列匹配
# ============================================================
@app.route('/api/match', methods=['POST'])
def api_match():
    """执行AI列名匹配，支持多源文件"""
    data = get_session()
    
    if not data['source_files']:
        return jsonify({'success': False, 'error': '请先上传原始表格'})
    if not data['reference_path']:
        return jsonify({'success': False, 'error': '请先上传参考模板'})
    
    use_ai = request.json.get('useAI', True) if request.json else True
    
    try:
        template = read_reference_template(data['reference_path'])
        target_headers = template['headers']
        data['reference_headers'] = target_headers
        
        # 从参考模板获取公式列名
        formula_headers = get_formula_headers(template)
        
        all_mappings_list = []
        combined_mappings = []   # 合并去重后的映射（UI展示用）
        seen_mappings = set()    # (source, target) 去重
        
        # 读取参考模板数据用于AI列采样
        ref_df = None
        try:
            ref_df = pd.read_excel(data['reference_path'], header=0, dtype=object)
            ref_df = ref_df.dropna(how='all').reset_index(drop=True)
        except Exception:
            pass
        
        for src_info in data['source_files']:
            df, source_headers, _ = read_source_file(src_info['path'])
            
            mappings = match_headers(
                source_headers=source_headers,
                target_headers=target_headers,
                use_ai=use_ai,
                source_df=df,
                reference_df=ref_df,
                reference_headers=target_headers,
                formula_columns=formula_headers,
            )
            
            # 构造可编辑映射
            editable = []
            mapped_sources = set()
            for m in mappings:
                editable.append({
                    'source': m['source_column'],
                    'target': m['target_column'],
                    'confidence': m.get('confidence', 'medium'),
                    'reason': m.get('reason', ''),
                })
                mapped_sources.add(m['source_column'])
                # 合并去重
                key = (m['source_column'], m['target_column'])
                if key not in seen_mappings:
                    seen_mappings.add(key)
                    combined_mappings.append(editable[-1])
            
            unmapped = [h for h in source_headers if h not in mapped_sources]
            
            all_mappings_list.append({
                'source_name': src_info['name'],
                'mappings': editable,
                'unmapped': unmapped,
            })
        
        data['mappings_list'] = all_mappings_list
        
        # 保存原始自动生成的映射（用于后续对比用户是否手动修改）
        data['original_mappings'] = [
            {'source': m['source'], 'target': m['target']}
            for m in combined_mappings
        ]
        
        # 收集所有源文件的所有列名（去重）
        all_source_headers = []
        seen_headers = set()
        for src_info in data['source_files']:
            for h in src_info['headers']:
                if h not in seen_headers:
                    seen_headers.add(h)
                    all_source_headers.append(h)
        
        # 未映射的源列（跨所有文件去重）
        all_unmapped = []
        seen_unmapped = set()
        for item in all_mappings_list:
            for h in item['unmapped']:
                if h not in seen_unmapped:
                    seen_unmapped.add(h)
                    all_unmapped.append(h)
        
        return jsonify({
            'success': True,
            'sourceFiles': [{'name': f['name'], 'rows': f['rows']} for f in data['source_files']],
            'mappings': combined_mappings,
            'unmappedSources': all_unmapped,
            'targetOptions': target_headers,
            'sourceHeaders': all_source_headers,
            'allMappingsList': all_mappings_list,  # 每个文件的详细映射
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': f'匹配失败: {str(e)}'})


# ============================================================
# API 路由 - 执行转换
# ============================================================
@app.route('/api/convert', methods=['POST'])
def api_convert():
    """执行表格转换（支持多源文件批量处理并合并）"""
    data = get_session()
    
    if not data['source_files']:
        return jsonify({'success': False, 'error': '请先上传原始表格'})
    if not data['reference_path']:
        return jsonify({'success': False, 'error': '请先上传参考模板'})
    
    # 获取用户调整过的映射
    user_mappings = None
    if request.json and 'mappings' in request.json:
        user_mappings = request.json['mappings']
    
    try:
        source_files = data['source_files']
        
        if len(source_files) == 1:
            # 单文件：直接使用现有流程
            src = source_files[0]
            output_filename = f"{src['name']}_转换结果.xlsx"
            output_path = os.path.join(OUTPUT_DIR, output_filename)
            
            # 转换映射格式（前端用 source/target，后端用 source_column/target_column）
            converted_mappings = []
            if user_mappings:
                converted_mappings = [{
                    'source_column': m.get('source', m.get('source_column', '')),
                    'target_column': m.get('target', m.get('target_column', '')),
                } for m in user_mappings if m.get('target') or m.get('target_column')]
            
            result_path = _do_convert(
                source_path=src['path'],
                reference_path=data['reference_path'],
                output_path=output_path,
                mappings=converted_mappings,
                source_name=src['name'],
            )
        else:
            # 多文件：逐个转换后合并
            output_filename = f"批量转换结果_{len(source_files)}个文件.xlsx"
            output_path = os.path.join(OUTPUT_DIR, output_filename)
            result_path = _do_batch_convert(
                source_files=source_files,
                reference_path=data['reference_path'],
                output_path=output_path,
                user_mappings=user_mappings,
            )
        
        data['output_path'] = result_path
        data['output_name'] = os.path.basename(result_path)
        
        # 保存用户手动调整的映射到历史记录
        if user_mappings:
            original_mappings = data.get('original_mappings', [])
            # 收集所有源表的列名
            all_source_headers = []
            for src_info in data['source_files']:
                all_source_headers.extend(src_info['headers'])
            try:
                batch_add_mappings(
                    mappings=user_mappings,
                    source_headers=all_source_headers,
                    original_mappings=original_mappings,
                )
                print(f"  📝 已保存用户映射调整到历史记录")
            except Exception as e:
                print(f"  ⚠️ 保存映射历史失败: {e}")
        
        # 生成预览（取第一个源文件与输出对比）
        first_source = source_files[0]['path']
        comparison = _build_comparison(first_source, result_path)
        
        return jsonify({
            'success': True,
            'outputName': os.path.basename(result_path),
            'downloadUrl': '/api/download',
            'comparison': comparison,
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': f'转换失败: {str(e)}'})


def _build_comparison(source_path, output_path):
    """构建源表与结果表的对比预览数据（全列，最多20行）"""
    try:
        import openpyxl
        
        # 读取源表
        src_df, src_headers, _ = read_source_file(source_path)
        src_preview = []
        max_rows = min(20, len(src_df))
        for i in range(max_rows):
            row = {}
            for h in src_headers:
                val = src_df.iloc[i][h] if h in src_df.columns else ''
                if hasattr(val, 'item'):
                    val = val.item()
                if val is not None and str(val) != 'nan':
                    row[str(h)] = str(val)
                else:
                    row[str(h)] = ''
            src_preview.append(row)
        
        # 读取输出表
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
        out_headers = []
        for c in range(1, ws.max_column + 1):
            h = ws.cell(row=1, column=c).value
            out_headers.append(str(h) if h else '')
        
        out_preview = []
        out_max_rows = min(20, ws.max_row - 1)  # minus header row
        for r in range(2, out_max_rows + 2):
            row = {}
            for c in range(1, ws.max_column + 1):
                h = out_headers[c - 1] if c - 1 < len(out_headers) else f'Col{c}'
                val = ws.cell(row=r, column=c).value
                row[str(h)] = str(val) if val is not None else ''
            out_preview.append(row)
        wb.close()
        
        return {
            'sourceHeaders': src_headers,
            'sourcePreview': src_preview,
            'sourceRowCount': len(src_df),
            'outputHeaders': out_headers,
            'outputPreview': out_preview,
            'outputRowCount': ws.max_row - 1,
        }
    except Exception as e:
        return {'error': str(e)}


def _do_convert(source_path, reference_path, output_path, mappings, source_name):
    """执行实际的转换，使用已保存的映射"""
    import pandas as pd
    import openpyxl
    from openpyxl.utils import get_column_letter
    from copy import copy
    import re
    
    # 读取数据
    df, source_headers, source_header_row = read_source_file(source_path)
    template = read_reference_template(reference_path)
    target_headers = template['headers']
    formula_info = template['formula_info']
    fixed_values = template['fixed_values']
    formula_deps = template['formula_dependencies']
    
    num_target_cols = template['max_column']
    
    # 动态检测毛利率列（含"毛利率"的表头列）
    margin_columns = []
    for idx, h in enumerate(target_headers):
        if '毛利率' in normalize_header(h):
            margin_columns.append(get_column_letter(idx + 1))
    
    # 构建列映射：源列名 → 目标列索引(0-based)
    column_mapping = {}
    
    # 检查是否格式相同
    same_format = (
        len(source_headers) == len(target_headers) and
        sum(1 for sh in source_headers if any(
            normalize_header(sh) == normalize_header(th) for th in target_headers
        )) >= len(source_headers) * 0.7
    )
    
    if same_format:
        for i, sh in enumerate(source_headers):
            if i < len(target_headers):
                column_mapping[sh] = i
    else:
        for m in mappings:
            src_col = m['source_column']
            tgt_col = m['target_column']
            
            src_idx = None
            for i, h in enumerate(source_headers):
                if normalize_header(h) == normalize_header(src_col):
                    src_idx = i
                    break
            
            tgt_idx = None
            tgt_clean = normalize_header(tgt_col)
            for idx, h in enumerate(target_headers):
                if normalize_header(h) == tgt_clean:
                    tgt_idx = idx
                    break
            if tgt_idx is None:
                for idx, h in enumerate(target_headers):
                    h_clean = normalize_header(h)
                    if tgt_clean in h_clean or h_clean in tgt_clean:
                        tgt_idx = idx
                        break
            
            if src_idx is not None and tgt_idx is not None:
                column_mapping[src_col] = tgt_idx
    
    # 创建输出工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    # 写入表头
    for col_idx in range(1, num_target_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = target_headers[col_idx - 1]
        col_letter = get_column_letter(col_idx)
        if col_letter in template['header_styles']:
            styles = template['header_styles'][col_letter]
            if styles.get('font'): cell.font = styles['font']
            if styles.get('alignment'): cell.alignment = styles['alignment']
            if styles.get('fill'): cell.fill = styles['fill']
            if styles.get('border'): cell.border = styles['border']
    
    # 写入数据
    for src_row_idx in range(len(df)):
        output_row = src_row_idx + 2
        
        # 第一遍：源数据
        for col_idx in range(1, num_target_cols + 1):
            cell = ws.cell(row=output_row, column=col_idx)
            col_letter = get_column_letter(col_idx)
            
            if col_letter in template['data_styles']:
                styles = template['data_styles'][col_letter]
                if styles.get('font'): cell.font = styles['font']
                if styles.get('alignment'): cell.alignment = styles['alignment']
                if styles.get('fill'): cell.fill = styles['fill']
                if styles.get('border'): cell.border = styles['border']
                if styles.get('number_format'): cell.number_format = styles['number_format']
            
            src_col_idx = None
            for src_col_name, tgt_idx in column_mapping.items():
                if tgt_idx == col_idx - 1:
                    for i, h in enumerate(source_headers):
                        if normalize_header(h) == normalize_header(src_col_name):
                            src_col_idx = i
                            break
                    break
            
            if src_col_idx is not None and src_col_idx < len(df.columns):
                val = df.iloc[src_row_idx, src_col_idx]
                if pd.notna(val):
                    val_str = str(val)
                    if val_str.startswith('=DISPIMG'):
                        cell.value = None
                    else:
                        try:
                            if isinstance(val, str) and val_str.strip() == '':
                                cell.value = None
                            elif isinstance(val, (int, float)):
                                if isinstance(val, float):
                                    rounded = round(val, 6)
                                    cell.value = int(rounded) if rounded == int(rounded) else rounded
                                else:
                                    cell.value = val
                            else:
                                cleaned = val_str.replace(',', '').strip()
                                if cleaned.replace('.', '').replace('-', '').isdigit():
                                    num_val = float(cleaned)
                                    rounded = round(num_val, 6)
                                    cell.value = int(rounded) if rounded == int(rounded) else rounded
                                else:
                                    cell.value = val
                        except (ValueError, TypeError):
                            cell.value = val
        
        # 第二遍：公式
        for col_idx in range(1, num_target_cols + 1):
            col_letter = get_column_letter(col_idx)
            if col_letter in formula_info:
                cell = ws.cell(row=output_row, column=col_idx)
                formula = formula_info[col_letter]
                
                def replace_row(match):
                    cl = match.group(1)
                    rn = int(match.group(2))
                    return f"{cl}{output_row}" if rn == 2 else match.group(0)
                
                adjusted = re.sub(r'([A-Z]+)(\d+)', replace_row, formula)
                cell.value = adjusted
        
        # 第三遍：写入统一毛利率配置值
        # 所有毛利率列统一使用配置的 MARGIN_RATE（默认10%），覆盖公式
        for col_idx in range(1, num_target_cols + 1):
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=output_row, column=col_idx)
            # 毛利率列：强制写入统一配置值
            if col_letter in margin_columns:
                cell.value = MARGIN_RATE
                continue
            if cell.value is not None:
                continue
            # 其他被公式引用的固定值常量
            if col_letter in formula_deps and col_letter in fixed_values:
                val = fixed_values[col_letter]
                if isinstance(val, (int, float)) and 0 < val < 1:
                    cell.value = val
    
    # 插入图片：从源文件提取嵌入图片并放入输出文件的图片列
    from openpyxl.drawing.image import Image as OpenpyxlImage
    import io as _io
    images_map = extract_images_from_source(source_path, header_row=source_header_row)
    target_img_col = None
    source_img_col = None
    
    # 方式A：从 column_mapping 中查找（AI/规则匹配结果）
    for src_col_name, tgt_idx in column_mapping.items():
        if '图片' in normalize_header(src_col_name):
            for i, h in enumerate(source_headers):
                if normalize_header(h) == normalize_header(src_col_name):
                    source_img_col = i
                    break
            target_img_col = tgt_idx
            break
    
    # 方式B：如果 column_mapping 中没有，直接在源表和目标表头中按列名查找
    if source_img_col is None:
        for i, h in enumerate(source_headers):
            if '图片' in normalize_header(h):
                source_img_col = i
                break
    if target_img_col is None and source_img_col is not None:
        for i, h in enumerate(target_headers):
            if '图片' in normalize_header(h):
                target_img_col = i
                break
    
    if images_map and source_img_col is not None and target_img_col is not None:
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        
        # 列宽 EMU
        target_col_letter = get_column_letter(target_img_col + 1)
        col_width_chars = template['column_widths'].get(target_col_letter, None)
        if col_width_chars:
            col_width_emu = int(col_width_chars * 7 * 914400 / 96)
        else:
            col_width_emu = 1143000
        
        # 行高与边距（从 config 读取）
        ROW_HEIGHT_PT = IMAGE_ROW_HEIGHT_PT
        MARGIN_PT = IMAGE_MARGIN_PT
        MAX_IMG_H_EMU = int((ROW_HEIGHT_PT - MARGIN_PT * 2) * 12700)
        ROW_HEIGHT_EMU = int(ROW_HEIGHT_PT * 12700)
        
        # 统一设置所有数据行高度为 60pt
        for r in range(2, len(df) + 2):
            ws.row_dimensions[r].height = ROW_HEIGHT_PT
        
        # 计算每张图片缩放后尺寸（按高度限制等比缩放）
        img_sizes = {}
        for (src_row, src_col), img_data in images_map.items():
            if src_col == source_img_col and src_row < len(df):
                try:
                    test_img = OpenpyxlImage(_io.BytesIO(img_data))
                    orig_w = max(test_img.width, 1)
                    orig_h = max(test_img.height, 1)
                    orig_h_emu = int(orig_h * 914400 / 96)
                    orig_w_emu = int(orig_w * 914400 / 96)
                    
                    if orig_h_emu > MAX_IMG_H_EMU:
                        ratio = MAX_IMG_H_EMU / orig_h_emu
                        scaled_h_emu = MAX_IMG_H_EMU
                        scaled_w_emu = int(orig_w_emu * ratio)
                    else:
                        scaled_h_emu = orig_h_emu
                        scaled_w_emu = orig_w_emu
                    
                    # 水平居中偏移
                    if scaled_w_emu < col_width_emu:
                        col_off = (col_width_emu - scaled_w_emu) // 2
                    else:
                        col_off = 0
                    
                    img_sizes[(src_row, src_col)] = (scaled_w_emu, scaled_h_emu, col_off)
                except Exception:
                    pass
        
        # 插入图片（居中）
        for (src_row, src_col), img_data in images_map.items():
            if src_col == source_img_col and src_row < len(df):
                size = img_sizes.get((src_row, src_col))
                if size is None:
                    continue
                emu_w, emu_h, col_off = size
                try:
                    img_stream = _io.BytesIO(img_data)
                    img = OpenpyxlImage(img_stream)
                    row_off = (ROW_HEIGHT_EMU - emu_h) // 2
                    
                    from_marker = AnchorMarker(
                        col=target_img_col, colOff=col_off,
                        row=src_row + 1, rowOff=row_off
                    )
                    anchor = OneCellAnchor(
                        _from=from_marker,
                        ext=XDRPositiveSize2D(cx=emu_w, cy=emu_h)
                    )
                    ws.add_image(img, anchor)
                except Exception:
                    pass
    
    # 列宽
    for col_idx in range(1, num_target_cols + 1):
        col_letter = get_column_letter(col_idx)
        if col_letter in template['column_widths'] and template['column_widths'][col_letter]:
            ws.column_dimensions[col_letter].width = template['column_widths'][col_letter]
    
    wb.save(output_path)
    wb.close()
    
    return output_path


def _do_batch_convert(source_files, reference_path, output_path, user_mappings=None):
    """批量转换多个源文件并合并到一个输出文件（数据+图片）"""
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from copy import copy
    import io as _io
    
    template = read_reference_template(reference_path)
    target_headers = template['headers']
    num_target_cols = template['max_column']
    
    # 查找图片目标列
    target_img_col = None
    for i, h in enumerate(target_headers):
        if '图片' in normalize_header(h):
            target_img_col = i
            break
    
    # 列宽参数
    if target_img_col is not None:
        tgt_letter = get_column_letter(target_img_col + 1)
        col_width_chars = template['column_widths'].get(tgt_letter, None)
        col_width_emu = int(col_width_chars * 7 * 914400 / 96) if col_width_chars else 1143000
    else:
        col_width_emu = 1143000
    
    ROW_HEIGHT_PT = IMAGE_ROW_HEIGHT_PT
    MARGIN_PT = IMAGE_MARGIN_PT
    MAX_IMG_H_EMU = int((ROW_HEIGHT_PT - MARGIN_PT * 2) * 12700)
    ROW_HEIGHT_EMU = int(ROW_HEIGHT_PT * 12700)
    
    # 创建输出工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    # 写入表头（只写一次）
    for col_idx in range(1, num_target_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = target_headers[col_idx - 1]
        col_letter = get_column_letter(col_idx)
        if col_letter in template['header_styles']:
            styles = template['header_styles'][col_letter]
            if styles.get('font'): cell.font = styles['font']
            if styles.get('alignment'): cell.alignment = styles['alignment']
            if styles.get('fill'): cell.fill = styles['fill']
            if styles.get('border'): cell.border = styles['border']
    
    current_output_row = 2  # 数据从第2行开始
    
    # 逐个处理每个源文件
    for file_idx, src_info in enumerate(source_files):
        source_path = src_info['path']
        source_name = src_info['name']
        
        # 生成临时转换文件（含数据和图片）
        temp_output = os.path.join(OUTPUT_DIR, f"_temp_batch_{file_idx}.xlsx")
        
        # 确定映射
        if user_mappings:
            mappings = [{
                'source_column': m.get('source', m.get('source_column', '')),
                'target_column': m.get('target', m.get('target_column', '')),
            } for m in user_mappings if m.get('target') or m.get('target_column')]
        else:
            df, src_headers, _ = read_source_file(source_path)
            # 读取参考模板数据用于AI列采样
            ref_df = None
            try:
                ref_df = pd.read_excel(reference_path, header=0, dtype=object)
                ref_df = ref_df.dropna(how='all').reset_index(drop=True)
            except Exception:
                pass
            mappings_raw = match_headers(
                src_headers, target_headers,
                use_ai=True,
                source_df=df,
                reference_df=ref_df,
                reference_headers=target_headers,
                formula_columns=get_formula_headers(template),
            )
            mappings = [{
                'source_column': m['source_column'],
                'target_column': m['target_column'],
            } for m in mappings_raw]
        
        _do_convert(source_path, reference_path, temp_output, mappings, source_name)
        
        # 读取临时文件，复制数据到主表
        temp_wb = openpyxl.load_workbook(temp_output)
        temp_ws = temp_wb.active
        
        rows_in_this_file = temp_ws.max_row - 1  # 减去表头
        
        for row_idx in range(2, temp_ws.max_row + 1):
            for col_idx in range(1, num_target_cols + 1):
                src_cell = temp_ws.cell(row=row_idx, column=col_idx)
                dst_cell = ws.cell(row=current_output_row, column=col_idx)
                if src_cell.value is not None:
                    dst_cell.value = src_cell.value
                if src_cell.font: dst_cell.font = copy(src_cell.font)
                if src_cell.alignment: dst_cell.alignment = copy(src_cell.alignment)
                if src_cell.fill and src_cell.fill.start_color and src_cell.fill.start_color.rgb != '00000000':
                    dst_cell.fill = copy(src_cell.fill)
                if src_cell.border: dst_cell.border = copy(src_cell.border)
                if src_cell.number_format and src_cell.number_format != 'General':
                    dst_cell.number_format = src_cell.number_format
            
            if row_idx in temp_ws.row_dimensions:
                h = temp_ws.row_dimensions[row_idx].height
                if h:
                    ws.row_dimensions[current_output_row].height = h
            
            current_output_row += 1
        
        # 从源文件直接提取图片并插入主表
        source_img_col = None
        df_src, src_headers, src_hr = read_source_file(source_path)
        for i, h in enumerate(src_headers):
            if '图片' in normalize_header(h):
                source_img_col = i
                break
        
        if target_img_col is not None and source_img_col is not None:
            images_map = extract_images_from_source(source_path, header_row=src_hr)
            # 计算该文件在主表中的起始行偏移
            row_offset = current_output_row - rows_in_this_file - 2
            
            for (src_row, src_col), img_data in images_map.items():
                if src_col == source_img_col and src_row < len(df_src):
                    try:
                        test_img = XLImage(_io.BytesIO(img_data))
                        orig_w = max(test_img.width, 1)
                        orig_h = max(test_img.height, 1)
                        orig_h_emu = int(orig_h * 914400 / 96)
                        orig_w_emu = int(orig_w * 914400 / 96)
                        
                        if orig_h_emu > MAX_IMG_H_EMU:
                            ratio = MAX_IMG_H_EMU / orig_h_emu
                            scaled_h = MAX_IMG_H_EMU
                            scaled_w = int(orig_w_emu * ratio)
                        else:
                            scaled_h = orig_h_emu
                            scaled_w = orig_w_emu
                        
                        col_off = (col_width_emu - scaled_w) // 2 if scaled_w < col_width_emu else 0
                        row_off = (ROW_HEIGHT_EMU - scaled_h) // 2
                        
                        target_row = src_row + 2 + row_offset
                        img = XLImage(_io.BytesIO(img_data))
                        from_marker = AnchorMarker(
                            col=target_img_col, colOff=col_off,
                            row=target_row - 1, rowOff=row_off
                        )
                        anchor = OneCellAnchor(
                            _from=from_marker,
                            ext=XDRPositiveSize2D(cx=scaled_w, cy=scaled_h)
                        )
                        ws.add_image(img, anchor)
                    except Exception:
                        pass
        
        temp_wb.close()
        try:
            os.remove(temp_output)
        except Exception:
            pass
    
    # 列宽
    for col_idx in range(1, num_target_cols + 1):
        col_letter = get_column_letter(col_idx)
        if col_letter in template['column_widths'] and template['column_widths'][col_letter]:
            ws.column_dimensions[col_letter].width = template['column_widths'][col_letter]
    
    wb.save(output_path)
    wb.close()
    return output_path


# ============================================================
# API 路由 - 文件下载
# ============================================================
@app.route('/api/download')
def download_file():
    """下载转换后的文件（从会话中获取路径，避免中文文件名URL编码问题）"""
    data = get_session()
    filepath = data.get('output_path', '')
    
    if not filepath or not os.path.exists(filepath):
        return jsonify({'success': False, 'error': '文件不存在，请先执行转换'}), 404
    
    download_name = data.get('output_name', 'result.xlsx')
    
    return send_file(
        filepath,
        as_attachment=True,
        download_name=download_name,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


# ============================================================
# API 路由 - 获取会话状态
# ============================================================
@app.route('/api/status')
def api_status():
    """获取当前会话状态"""
    data = get_session()
    return jsonify({
        'hasSource': len(data.get('source_files', [])) > 0,
        'sourceFiles': [{'name': f['name'], 'rows': f['rows']} for f in data.get('source_files', [])],
        'hasReference': data['reference_path'] is not None,
        'referenceHeaders': data.get('reference_headers', []),
        'formulaCount': data.get('formula_count', 0),
        'hasMappings': len(data.get('mappings_list', [])) > 0,
        'mappingCount': sum(len(m['mappings']) for m in data.get('mappings_list', [])),
        'hasOutput': data.get('output_path') is not None,
        'outputName': data.get('output_name', ''),
    })


@app.route('/api/reset', methods=['POST'])
def api_reset():
    """重置会话"""
    sid = session.get('sid')
    if sid in sessions:
        data = sessions[sid]
        for key in ('source_paths', 'source_files'):
            items = data.get(key, [])
            if isinstance(items, list):
                for item in items:
                    p = item if isinstance(item, str) else item.get('path', '')
                    if p and os.path.exists(p):
                        try: os.remove(p)
                        except: pass
        if data.get('reference_path') and os.path.exists(data['reference_path']):
            try: os.remove(data['reference_path'])
            except: pass
        # 清理旧的单文件字段
        for old_key in ('source_path',):
            if data.get(old_key) and os.path.exists(data[old_key]):
                try: os.remove(data[old_key])
                except: pass
        del sessions[sid]
    
    session.pop('sid', None)
    return jsonify({'success': True})


# ============================================================
# 启动入口
# ============================================================
if __name__ == '__main__':
    import traceback as _tb
    print("\n" + "=" * 60)
    print("  📊 表格转换工具 - Web 界面")
    print("=" * 60)
    print(f"  🌐 访问地址: http://127.0.0.1:5000")
    print(f"  📂 上传目录: {UPLOAD_DIR}")
    print(f"  📁 输出目录: {OUTPUT_DIR}")
    cleanup_old_files()
    print("=" * 60 + "\n")
    
    try:
        app.run(debug=False, host='0.0.0.0', port=5000, threaded=True)
    except Exception as _e:
        print(f"\n❌ 服务异常退出: {_e}", flush=True)
        _tb.print_exc()
